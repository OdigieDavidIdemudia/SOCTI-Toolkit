import csv
import argparse
import sys
import os
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

IP_ALIASES = ["ip", "ipaddress", "ip address", "ip_address"]
HOST_ALIASES = ["hostname", "host name", "host", "device name"]

# Default config fallback
CONFIG = {
    "jobs": [
        {"label": "Cortex (Servers & VPCs)", "sheet_name": "Cortex Servers-VPCs", "tool_file": None, "inventory_file": None},
        {"label": "Cortex (ATMs)", "sheet_name": "Cortex ATMs", "tool_file": None, "inventory_file": None},
        {"label": "SIEM", "sheet_name": "SIEM", "tool_file": None, "inventory_file": None},
        {"label": "DAM", "sheet_name": "DAM", "tool_file": None, "inventory_file": None}
    ],
    "output_path": "Asset_Compliance_Report.xlsx"
}

def detect_delimiter(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            first_line = f.readline()
            if ',' in first_line:
                return ','
            elif '\t' in first_line:
                return '\t'
            elif ';' in first_line:
                return ';'
            return ','
    except Exception:
        return ','

def parse_file(file_path):
    """
    Parses a CSV/TSV file and returns a dictionary mapping IP -> Hostname (if found).
    """
    if not file_path or not os.path.exists(file_path):
        return {}

    delim = detect_delimiter(file_path)
    data = {}
    
    with open(file_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f, delimiter=delim)
        try:
            headers = next(reader)
        except StopIteration:
            return {}

        headers_lower = [h.strip().lower() for h in headers]
        
        ip_idx = -1
        host_idx = -1

        for i, h in enumerate(headers_lower):
            if ip_idx == -1 and h in IP_ALIASES:
                ip_idx = i
            if host_idx == -1 and h in HOST_ALIASES:
                host_idx = i

        if ip_idx == -1:
            raise ValueError(f"Could not find an IP address column in {file_path}. Header found: {headers}")

        for row in reader:
            if len(row) <= ip_idx:
                continue
            
            ip_raw = row[ip_idx].strip().lower()
            if not ip_raw:
                continue
            
            host = ""
            if host_idx != -1 and len(row) > host_idx:
                host = row[host_idx].strip()

            ips = [i.strip() for i in re.split(r'[,\s;]+', ip_raw) if i.strip()]
            
            for single_ip in ips:
                if single_ip not in data:
                    data[single_ip] = host

    return data

def run_compliance_check(jobs, output_path):
    valid_jobs = []
    for job in jobs:
        if job.get('tool_file') and job.get('inventory_file'):
            valid_jobs.append(job)

    if not valid_jobs:
        print("No complete tool/inventory pairs were provided - nothing to do.")
        sys.exit(1)

    wb = Workbook()
    # Remove default sheet later
    default_sheet = wb.active

    # Styles
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")
    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")

    summary_data = []

    for job in valid_jobs:
        try:
            tool_data = parse_file(job['tool_file'])
            inventory_data = parse_file(job['inventory_file'])
        except Exception as e:
            print(f"Error processing job '{job['label']}': {e}")
            continue

        compliant = []
        not_compliant = []

        for inv_ip, inv_host in inventory_data.items():
            if inv_ip in tool_data:
                compliant.append((inv_host, inv_ip))
            else:
                not_compliant.append((inv_host, inv_ip))

        compliant.sort(key=lambda x: x[1])
        not_compliant.sort(key=lambda x: x[1])

        sheet_name = job['sheet_name'][:31] # Excel limit
        ws = wb.create_sheet(title=sheet_name)

        # Headers
        ws.merge_cells('B1:C1')
        ws['B1'] = 'Not Compliant'
        ws['B1'].font = bold_font
        ws['B1'].alignment = center_align
        ws['B1'].fill = red_fill

        ws.merge_cells('E1:F1')
        ws['E1'] = 'Compliant'
        ws['E1'].font = bold_font
        ws['E1'].alignment = center_align
        ws['E1'].fill = green_fill

        ws['B2'] = 'Hostname'
        ws['C2'] = 'IP address'
        ws['E2'] = 'Hostname'
        ws['F2'] = 'IP address'

        for cell in ['B2', 'C2', 'E2', 'F2']:
            ws[cell].font = bold_font
            ws[cell].alignment = center_align

        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20

        max_rows = max(len(not_compliant), len(compliant))
        for i in range(max_rows):
            row_idx = i + 3
            if i < len(not_compliant):
                ws.cell(row=row_idx, column=2, value=not_compliant[i][0])
                ws.cell(row=row_idx, column=3, value=not_compliant[i][1])
            if i < len(compliant):
                ws.cell(row=row_idx, column=5, value=compliant[i][0])
                ws.cell(row=row_idx, column=6, value=compliant[i][1])

        summary_data.append({
            "label": job['label'],
            "sheet_name": sheet_name
        })

    if summary_data:
        ws_sum = wb.create_sheet(title="Summary", index=0)
        
        headers = ["Tool", "Compliant", "Not Compliant", "Total", "% Compliant"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws_sum.cell(row=3, column=col_idx, value=header)
            cell.font = bold_font

        for row_idx, job_info in enumerate(summary_data, start=4):
            sheet = job_info['sheet_name']
            ws_sum.cell(row=row_idx, column=1, value=job_info['label'])
            
            # Formulas
            ws_sum.cell(row=row_idx, column=2, value=f"=COUNTA('{sheet}'!F3:F100000)")
            ws_sum.cell(row=row_idx, column=3, value=f"=COUNTA('{sheet}'!C3:C100000)")
            ws_sum.cell(row=row_idx, column=4, value=f"=B{row_idx}+C{row_idx}")
            pct_cell = ws_sum.cell(row=row_idx, column=5, value=f"=IF(D{row_idx}=0,0,B{row_idx}/D{row_idx})")
            pct_cell.number_format = '0%'

        for col in ['A', 'B', 'C', 'D', 'E']:
            ws_sum.column_dimensions[col].width = 18

    # Clean up empty default sheet
    if default_sheet.title == "Sheet":
        wb.remove(default_sheet)

    wb.save(output_path)
    print(f"Report generated successfully at {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Asset Compliance Checker")
    parser.add_argument('--siem-tool', help='Path to SIEM tool export')
    parser.add_argument('--siem-inventory', help='Path to SIEM inventory file')
    parser.add_argument('--cortex-servers-tool', help='Path to Cortex Servers tool export')
    parser.add_argument('--cortex-servers-inventory', help='Path to Cortex Servers inventory file')
    parser.add_argument('--cortex-atms-tool', help='Path to Cortex ATMs tool export')
    parser.add_argument('--cortex-atms-inventory', help='Path to Cortex ATMs inventory file')
    parser.add_argument('--dam-tool', help='Path to DAM tool export')
    parser.add_argument('--dam-inventory', help='Path to DAM inventory file')
    parser.add_argument('--output', default='Asset_Compliance_Report.xlsx', help='Output Excel file path')

    args = parser.parse_args()

    # Determine if CLI args were passed
    any_args_passed = any([
        args.siem_tool, args.siem_inventory,
        args.cortex_servers_tool, args.cortex_servers_inventory,
        args.cortex_atms_tool, args.cortex_atms_inventory,
        args.dam_tool, args.dam_inventory
    ])

    if any_args_passed:
        jobs = []
        if args.cortex_servers_tool and args.cortex_servers_inventory:
            jobs.append({"label": "Cortex (Servers & VPCs)", "sheet_name": "Cortex Servers-VPCs", "tool_file": args.cortex_servers_tool, "inventory_file": args.cortex_servers_inventory})
        if args.cortex_atms_tool and args.cortex_atms_inventory:
            jobs.append({"label": "Cortex (ATMs)", "sheet_name": "Cortex ATMs", "tool_file": args.cortex_atms_tool, "inventory_file": args.cortex_atms_inventory})
        if args.siem_tool and args.siem_inventory:
            jobs.append({"label": "SIEM", "sheet_name": "SIEM", "tool_file": args.siem_tool, "inventory_file": args.siem_inventory})
        if args.dam_tool and args.dam_inventory:
            jobs.append({"label": "DAM", "sheet_name": "DAM", "tool_file": args.dam_tool, "inventory_file": args.dam_inventory})
        output_path = args.output
    else:
        jobs = CONFIG['jobs']
        output_path = CONFIG['output_path']

    try:
        run_compliance_check(jobs, output_path)
    except SystemExit:
        raise
    except Exception as e:
        print(f"Failed to generate report: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()
