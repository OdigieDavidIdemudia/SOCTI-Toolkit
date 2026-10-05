from pathlib import Path
from typing import List, Union
from .chunker import Chunk

class DocumentIngestor:
    @staticmethod
    def ingest(file_path: Union[str, Path]) -> Union[List[Chunk], str]:
        """Reads a file and returns either a list of pre-classified Chunks or raw text for the heuristic chunker."""
        path = Path(file_path)
        ext = path.suffix.lower()
        
        if ext == '.txt':
            return path.read_text(encoding='utf-8', errors='ignore')
            
        elif ext == '.docx':
            return DocumentIngestor._ingest_docx(path)
            
        elif ext == '.pdf':
            return DocumentIngestor._ingest_pdf(path)
            
        elif ext == '.xlsx':
            return DocumentIngestor._ingest_xlsx(path)
            
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def _ingest_docx(path: Path) -> List[Chunk]:
        import docx
        doc = docx.Document(path)
        chunks = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            
            style_name = para.style.name.lower()
            hint = None
            if 'heading' in style_name:
                hint = 'heading'
            elif 'list' in style_name or 'bullet' in style_name:
                hint = 'bullet_list'
            elif 'code' in style_name:
                hint = 'code'
                
            chunks.append(Chunk(text=text, hint=hint))
            
        # We can also handle tables
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                table_text.append(row_data)
            
            if table_text and len(table_text) > 0:
                header = " | ".join(table_text[0])
                separator = " | ".join(["---"] * len(table.rows[0].cells))
                table_str = f"| {header} |\n| {separator} |\n"
                for row in table_text[1:]:
                    table_str += f"| {' | '.join(row)} |\n"
                chunks.append(Chunk(text=table_str, hint="table"))
                
        return chunks

    @staticmethod
    def _ingest_pdf(path: Path) -> str:
        import fitz  # PyMuPDF
        text_content = []
        with fitz.open(str(path)) as doc:
            for page in doc:
                extracted = page.get_text()
                if extracted.strip():
                    text_content.append(extracted.strip())
        return "\n\n".join(text_content)

    @staticmethod
    def _ingest_xlsx(path: Path) -> List[Chunk]:
        import pandas as pd
        # Read excel sheets
        dfs = pd.read_excel(path, sheet_name=None)
        chunks = []
        for sheet_name, df in dfs.items():
            chunks.append(Chunk(text=f"Sheet: {sheet_name}", hint="heading"))
            # convert dataframe to markdown table string
            # we need tabulate installed for pandas to_markdown, we will add it to requirements or just output csv.
            # let's write a simple markdown generator to avoid tabulate dependency if it's missing
            csv_str = df.to_csv(index=False)
            lines = csv_str.strip().split('\n')
            if not lines:
                continue
                
            header = lines[0].replace(',', ' | ')
            separator = " | ".join(["---"] * len(df.columns))
            table_str = f"| {header} |\n| {separator} |\n"
            for row in lines[1:]:
                table_str += f"| {row.replace(',', ' | ')} |\n"
            
            chunks.append(Chunk(text=table_str, hint="table"))
            
        return chunks
