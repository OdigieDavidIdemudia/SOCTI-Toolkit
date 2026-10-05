import argparse
import csv
import sys
import re
from typing import Dict, List, Tuple
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

class MissingColumnError(Exception):
    pass

def normalize_hostname(hostname: str) -> str:
    if not hostname:
        return ""
    return re.sub(r'[-_]', '', hostname.lower().strip())

def parse_report_file(filepath: str) -> List[Dict[str, str]]:
    """
    Parses the XLSX report and deduplicates based on (ip, hostname).
    """
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        print(f"Failed to load report file '{filepath}': {e}")
        sys.exit(1)
        
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Report file is empty")
        
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    
    ip_aliases = ["ip_address", "ip address", "ip", "ipaddress"]
    host_aliases = ["hostname", "host name", "host", "endpoint_name"]
    
    ip_col = None
    host_col = None
    
    for idx, header in enumerate(headers):
        if header in ip_aliases and ip_col is None:
            ip_col = idx
        if header in host_aliases and host_col is None:
            host_col = idx
            
    if ip_col is None or host_col is None:
        raise MissingColumnError(f"Report is missing IP/Hostname columns. Header found: {headers}")
        
    dedup = set()
    records = []
    
    for row in rows[1:]:
        ip = str(row[ip_col]).strip() if len(row) > ip_col and row[ip_col] else ""
        host = str(row[host_col]).strip() if len(row) > host_col and row[host_col] else ""
        
        if not ip and not host:
            continue
            
        key = (ip, host.lower())
        if key not in dedup:
            dedup.add(key)
            records.append({
                "ip_address": ip,
                "hostname": host
            })
            
    return records

def sniff_dialect(sample: str) -> csv.Dialect:
    try:
        return csv.Sniffer().sniff(sample, delimiters=[',', '\t', ';'])
    except csv.Error:
        csv.register_dialect('fallback', delimiter='\t')
        return csv.get_dialect('fallback')

def parse_tsv_file(filepath: str) -> List[Dict]:
    """
    Parses the TSV file and filters for DRIVE_FIXED.
    Groups by IP and normalized hostname.
    """
    try:
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            sample = f.read(4096)
            f.seek(0)
            dialect = sniff_dialect(sample)
            reader = csv.reader(f, dialect)
            rows = list(reader)
    except Exception as e:
        print(f"Failed to read TSV file '{filepath}': {e}")
        sys.exit(1)
        
    if not rows:
        raise ValueError("TSV file is empty")
        
    headers = [h.strip().lower() for h in rows[0]]
    required_cols = ["endpoint_name", "name", "drive_type", "ip_address"]
    
    col_idx = {}
    for req in required_cols:
        try:
            col_idx[req] = headers.index(req)
        except ValueError:
            raise MissingColumnError(f"TSV is missing expected columns. Found: {headers}, required: {required_cols}")
            
    groups = {}
    
    for row in rows[1:]:
        if not row: continue
        
        def safe_get(idx):
            return row[idx].strip() if len(row) > idx else ""
            
        endpoint_name = safe_get(col_idx["endpoint_name"])
        drive_name = safe_get(col_idx["name"])
        drive_type = safe_get(col_idx["drive_type"])
        ip_address = safe_get(col_idx["ip_address"])
        
        if drive_type.upper() != "DRIVE_FIXED":
            continue
            
        norm_host = normalize_hostname(endpoint_name)
        key = (ip_address, norm_host)
        
        if key not in groups:
            groups[key] = {
                "endpoint_name": endpoint_name,
                "ip_address": ip_address,
                "drives": set()
            }
        
        if drive_name:
            groups[key]["drives"].add(drive_name)
            
    return list(groups.values())

def build_indexes(tsv_groups: List[Dict]) -> Tuple[Dict, Dict, Dict]:
    by_ip = {}
    by_host_raw = {}
    by_host_norm = {}
    
    for group in tsv_groups:
        ip = group["ip_address"]
        host = group["endpoint_name"]
        
        if ip:
            by_ip[ip] = group
            
        if host:
            raw = host.lower().strip()
            norm = normalize_hostname(host)
            by_host_raw[raw] = group
            by_host_norm[norm] = group
            
    return by_ip, by_host_raw, by_host_norm

def run(report_path: str, tsv_path: str, output_path: str) -> Tuple[int, int, int]:
    report_records = parse_report_file(report_path)
    tsv_groups = parse_tsv_file(tsv_path)
    by_ip, by_host_raw, by_host_norm = build_indexes(tsv_groups)
    
    still_pending = []
    resolved = []
    not_found = []
    
    for record in report_records:
        ip = record["ip_address"]
        host = record["hostname"]
        
        matched_group = None
        matched_by = ""
        
        if ip and ip in by_ip:
            matched_group = by_ip[ip]
            matched_by = "IP address"
        elif host and host.lower().strip() in by_host_raw:
            matched_group = by_host_raw[host.lower().strip()]
            matched_by = "Hostname"
        elif host and normalize_hostname(host) in by_host_norm:
            matched_group = by_host_norm[normalize_hostname(host)]
            matched_by = "Hostname (normalized)"
            
        if matched_group:
            drives = matched_group["drives"]
            count = len(drives)
            drive_names = ", ".join(sorted(list(drives)))
            
            out_rec = {
                "Endpoint Name": matched_group["endpoint_name"],
                "Number of Drives": count,
                "Drive Names": drive_names,
                "IP Address": matched_group["ip_address"],
                "Matched By": matched_by
            }
            
            if count >= 2:
                still_pending.append(out_rec)
            else:
                resolved.append(out_rec)
        else:
            not_found.append({
                "Endpoint Name": host,
                "IP Address": ip
            })
            
    # Sort
    still_pending.sort(key=lambda x: x["Number of Drives"], reverse=True)
    resolved.sort(key=lambda x: str(x["Endpoint Name"]).lower())
    not_found.sort(key=lambda x: str(x["Endpoint Name"]).lower())
    
    wb = Workbook()
    
    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Category", "Count"])
    ws_summary.append(["Total tracked endpoints", len(report_records)])
    ws_summary.append(["Still Pending (2+ drives confirmed)", len(still_pending)])
    ws_summary.append(["Resolved (now 1 drive)", len(resolved)])
    ws_summary.append(["Not Found (not reporting)", len(not_found)])
    
    # Sheet 2: Still Pending
    ws_pending = wb.create_sheet(title="Still Pending")
    pending_headers = ["Endpoint Name", "Number of Drives", "Drive Names", "IP Address", "Matched By"]
    ws_pending.append(pending_headers)
    for row in still_pending:
        ws_pending.append([row[h] for h in pending_headers])
        
    # Sheet 3: Resolved
    ws_resolved = wb.create_sheet(title="Resolved")
    resolved_headers = ["Endpoint Name", "Number of Drives", "Drive Names", "IP Address", "Matched By"]
    ws_resolved.append(resolved_headers)
    for row in resolved:
        ws_resolved.append([row[h] for h in resolved_headers])
        
    # Sheet 4: Not Found
    ws_not_found = wb.create_sheet(title="Not Found")
    not_found_headers = ["Endpoint Name", "IP Address"]
    ws_not_found.append(not_found_headers)
    for row in not_found:
        ws_not_found.append([row[h] for h in not_found_headers])
        
    # Formatting
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            max_length = 0
            column = [c for c in col if c.value]
            if column:
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[get_column_letter(col[0].column)].width = adjusted_width
                
    wb.save(output_path)
    return len(still_pending), len(resolved), len(not_found)

def main():
    parser = argparse.ArgumentParser(description="Multi-Drive Pending Checker (Drive Audit Engine)")
    parser.add_argument("--report", required=True, help="Path to the tracked Excel report")
    parser.add_argument("--tsv", required=True, help="Path to the TSV/CSV tool export")
    parser.add_argument("--output", default="Multi_Drive_Status.xlsx", help="Path for the output Excel file")
    
    args = parser.parse_args()
    
    try:
        p, r, nf = run(args.report, args.tsv, args.output)
        print(f"Successfully generated report at {args.output}")
        print(f"Summary: {p} pending, {r} resolved, {nf} not found.")
    except Exception as e:
        print(f"Error: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()
